from io import BytesIO

import openpyxl
from django.test import override_settings
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APITestCase

from apps.core.models import Lab, User
from apps.core.test_helpers import RoleMatrixTestMixin
from apps.inventory.models import WarehouseItem


class WarehouseItemCrudApiTests(APITestCase):
    """Test CRUD operations for WarehouseItem model."""

    def setUp(self):
        self.lab_a = Lab.objects.create(name="Lab A")
        self.lab_b = Lab.objects.create(name="Lab B")

        self.admin_a = User.objects.create_user(
            username="admin_a",
            email="warehouse_admin_a@example.com",
            password="password123",
            role="admin",
            lab=self.lab_a,
        )
        self.user_a = User.objects.create_user(
            username="user_a",
            email="warehouse_user_a@example.com",
            password="password123",
            role="user",
            lab=self.lab_a,
        )
        self.admin_b = User.objects.create_user(
            username="admin_b",
            email="warehouse_admin_b@example.com",
            password="password123",
            role="admin",
            lab=self.lab_b,
        )
        self.superadmin = User.objects.create_user(
            username="warehouse_superadmin",
            email="warehouse_superadmin@example.com",
            password="password123",
            role="superadmin",
            is_superuser=True,
        )

    def test_create_warehouse_item(self):
        """Test creating a new warehouse item."""
        self.client.force_authenticate(user=self.admin_a)
        url = reverse("warehouseitem-list")
        payload = {
            "name": "Zircon blocks",
            "sku": "ZIR-001",
            "quantity": 50,
            "unit": "pcs",
            "min_threshold": 10,
            "category": "Materials",
            "location": "A1",
            "supplier": "Dental Supplier",
            "cost_price": 12.5,
            "notes": "High translucency",
        }

        response = self.client.post(url, payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["name"], "Zircon blocks")
        self.assertEqual(response.data["sku"], "ZIR-001")
        self.assertEqual(float(response.data["quantity"]), 50.0)
        self.assertEqual(response.data["supplier"], "Dental Supplier")
        self.assertEqual(response.data["lab"], self.lab_a.id)

    def test_superadmin_can_create_warehouse_item_for_selected_lab(self):
        self.client.force_authenticate(user=self.superadmin)
        response = self.client.post(
            reverse("warehouseitem-list"),
            {
                "lab": self.lab_b.id,
                "name": "Superadmin item",
                "sku": "SUPER-001",
                "quantity": 5,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["lab"], self.lab_b.id)

    def test_list_warehouse_items(self):
        """Test listing warehouse items scoped to lab."""
        WarehouseItem.objects.create(
            lab=self.lab_a,
            name="Item A1",
            sku="SKU-A1",
            quantity=10,
        )
        WarehouseItem.objects.create(
            lab=self.lab_a,
            name="Item A2",
            sku="SKU-A2",
            quantity=20,
        )
        WarehouseItem.objects.create(
            lab=self.lab_b,
            name="Item B1",
            sku="SKU-B1",
            quantity=30,
        )

        self.client.force_authenticate(user=self.admin_a)
        url = reverse("warehouseitem-list")

        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data), 2)  # Only lab_a items

    def test_get_warehouse_item(self):
        """Test retrieving a specific warehouse item."""
        item = WarehouseItem.objects.create(
            lab=self.lab_a,
            name="Test Item",
            sku="TEST-001",
            quantity=15,
        )

        self.client.force_authenticate(user=self.admin_a)
        url = reverse("warehouseitem-detail", args=[item.id])

        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["name"], "Test Item")

    def test_update_warehouse_item(self):
        """Test updating a warehouse item."""
        item = WarehouseItem.objects.create(
            lab=self.lab_a,
            name="Zircon blocks",
            sku="ZIR-001",
            quantity=50,
        )

        self.client.force_authenticate(user=self.admin_a)
        url = reverse("warehouseitem-detail", args=[item.id])
        payload = {
            "name": "Zircon blocks",
            "sku": "ZIR-001",
            "quantity": 42,
            "unit": "pcs",
        }

        response = self.client.put(url, payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(float(response.data["quantity"]), 42.0)

    def test_delete_warehouse_item(self):
        """Test deleting a warehouse item."""
        item = WarehouseItem.objects.create(
            lab=self.lab_a,
            name="Test Item",
            sku="DEL-001",
            quantity=10,
        )

        self.client.force_authenticate(user=self.admin_a)
        url = reverse("warehouseitem-detail", args=[item.id])

        response = self.client.delete(url)

        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)

        # Verify deletion
        response = self.client.get(url)
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_cross_lab_access_denied(self):
        """Test that users cannot access items from other labs."""
        item_b = WarehouseItem.objects.create(
            lab=self.lab_b,
            name="Lab B Item",
            sku="LB-001",
            quantity=5,
        )

        self.client.force_authenticate(user=self.admin_a)
        url = reverse("warehouseitem-detail", args=[item_b.id])

        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_non_admin_can_list_items(self):
        """Test that regular users in the lab can list warehouse items."""
        WarehouseItem.objects.create(
            lab=self.lab_a,
            name="Item A1",
            sku="SKU-A1",
            quantity=10,
        )

        self.client.force_authenticate(user=self.user_a)
        url = reverse("warehouseitem-list")

        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data), 1)

    def test_non_admin_cannot_create_warehouse_item(self):
        self.client.force_authenticate(user=self.user_a)
        response = self.client.post(
            reverse("warehouseitem-list"),
            {"name": "Blocked item", "quantity": 1},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertFalse(WarehouseItem.objects.filter(name="Blocked item").exists())

    def test_superadmin_can_list_items_across_labs(self):
        WarehouseItem.objects.create(
            lab=self.lab_a,
            name="Item A1",
            sku="SKU-A1",
            quantity=10,
        )
        WarehouseItem.objects.create(
            lab=self.lab_b,
            name="Item B1",
            sku="SKU-B1",
            quantity=20,
        )

        self.client.force_authenticate(user=self.superadmin)
        response = self.client.get(reverse("warehouseitem-list"))

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data), 2)

    def test_stats_are_scoped_to_authenticated_lab(self):
        WarehouseItem.objects.create(
            lab=self.lab_a,
            name="In stock",
            category="Materials",
            quantity="10.00",
            min_threshold="3.00",
            cost_price="2.50",
        )
        WarehouseItem.objects.create(
            lab=self.lab_a,
            name="Low stock",
            category="Materials",
            quantity="2.00",
            min_threshold="5.00",
            cost_price="10.00",
        )
        WarehouseItem.objects.create(
            lab=self.lab_a,
            name="Out of stock",
            category="Tools",
            quantity="0.00",
            min_threshold="1.00",
            cost_price="7.00",
        )
        WarehouseItem.objects.create(
            lab=self.lab_b,
            name="Other lab stock",
            category="Materials",
            quantity="100.00",
            min_threshold="5.00",
            cost_price="1.00",
        )

        self.client.force_authenticate(user=self.admin_a)
        response = self.client.get(reverse("warehouseitem-stats"))

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["total_items"], 3)
        self.assertEqual(response.data["total_value"], "45.00")
        self.assertEqual(response.data["low_stock_count"], 1)
        self.assertEqual(response.data["out_of_stock_count"], 1)
        self.assertEqual(
            response.data["categories"],
            [
                {"category": "Materials", "count": 2},
                {"category": "Tools", "count": 1},
            ],
        )

    def test_superadmin_stats_include_all_labs(self):
        WarehouseItem.objects.create(
            lab=self.lab_a,
            name="Item A",
            quantity="1.00",
            cost_price="5.00",
        )
        WarehouseItem.objects.create(
            lab=self.lab_b,
            name="Item B",
            quantity="2.00",
            cost_price="3.00",
        )

        self.client.force_authenticate(user=self.superadmin)
        response = self.client.get(reverse("warehouseitem-stats"))

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["total_items"], 2)
        self.assertEqual(response.data["total_value"], "11.00")


class WarehouseBulkImportTests(APITestCase):
    """Tests for POST /warehouse/import/ and /warehouse/import-partial/."""

    def setUp(self):
        self.lab = Lab.objects.create(name="Bulk Lab")
        self.admin = User.objects.create_user(
            username="bulk_admin",
            email="bulk_admin@example.com",
            password="password123",
            role="admin",
            lab=self.lab,
        )
        self.no_lab_user = User.objects.create_user(
            username="bulk_no_lab",
            email="bulk_no_lab@example.com",
            password="password123",
            role="user",
        )
        self.regular = User.objects.create_user(
            username="bulk_regular",
            email="bulk_regular@example.com",
            password="password123",
            role="user",
            lab=self.lab,
        )
        self.import_url = reverse("warehouseitem-bulk-import")
        self.import_partial_url = reverse("warehouseitem-bulk-import-partial")

    def _valid_items(self):
        return [
            {
                "name": "Item A",
                "sku": "SKU-A",
                "quantity": "10.00",
                "unit": "pcs",
                "supplier": "Supplier A",
                "cost_price": "5.50",
            },
            {"name": "Item B", "sku": "SKU-B", "quantity": "20.00", "unit": "pcs"},
        ]

    def test_bulk_import_happy_path(self):
        """All valid items are created atomically."""
        self.client.force_authenticate(user=self.admin)
        response = self.client.post(self.import_url, self._valid_items(), format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["imported"], 2)
        self.assertEqual(WarehouseItem.objects.filter(lab=self.lab).count(), 2)

    def test_bulk_import_validation_error_rolls_back(self):
        """Any invalid row causes the entire import to fail (no rows created)."""
        self.client.force_authenticate(user=self.admin)
        items = self._valid_items()
        items.append({"quantity": "bad"})  # missing required name

        response = self.client.post(self.import_url, items, format="json")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("errors", response.data)
        self.assertEqual(WarehouseItem.objects.filter(lab=self.lab).count(), 0)

    def test_bulk_import_lab_scoping(self):
        """Imported items are assigned to the authenticated user's lab."""
        self.client.force_authenticate(user=self.admin)
        self.client.post(self.import_url, self._valid_items(), format="json")

        items = WarehouseItem.objects.filter(lab=self.lab)
        self.assertEqual(items.count(), 2)
        for item in items:
            self.assertEqual(item.lab_id, self.lab.id)

    def test_bulk_import_no_lab_denied(self):
        """Users without a lab must receive 403."""
        self.client.force_authenticate(user=self.no_lab_user)
        response = self.client.post(self.import_url, self._valid_items(), format="json")
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_bulk_import_non_admin_denied(self):
        self.client.force_authenticate(user=self.regular)
        response = self.client.post(self.import_url, self._valid_items(), format="json")

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(WarehouseItem.objects.filter(lab=self.lab).count(), 0)

    def test_bulk_import_requires_list(self):
        """Sending a dict instead of a list returns 400."""
        self.client.force_authenticate(user=self.admin)
        response = self.client.post(self.import_url, {"name": "x"}, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_bulk_import_partial_skips_invalid(self):
        """Partial import skips invalid rows and imports the valid ones."""
        self.client.force_authenticate(user=self.admin)
        items = self._valid_items()
        items.append({"quantity": "bad"})  # invalid row

        response = self.client.post(self.import_partial_url, items, format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["imported"], 2)
        self.assertEqual(response.data["skipped"], 1)
        self.assertEqual(WarehouseItem.objects.filter(lab=self.lab).count(), 2)

    def test_bulk_import_decimal_precision(self):
        """Decimal fields are stored without float rounding artifacts."""
        self.client.force_authenticate(user=self.admin)
        items = [{"name": "Precise Item", "quantity": "1.10", "cost_price": "9.99"}]

        self.client.post(self.import_url, items, format="json")

        item = WarehouseItem.objects.get(lab=self.lab, name="Precise Item")
        self.assertEqual(str(item.quantity), "1.10")
        self.assertEqual(str(item.cost_price), "9.99")

    def test_bulk_import_supplier_field(self):
        self.client.force_authenticate(user=self.admin)
        items = [{"name": "Supplier Item", "supplier": "Dental Depot"}]

        response = self.client.post(self.import_url, items, format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        item = WarehouseItem.objects.get(lab=self.lab, name="Supplier Item")
        self.assertEqual(item.supplier, "Dental Depot")


class InventoryCSVImportTests(APITestCase):
    def setUp(self):
        self.lab = Lab.objects.create(name="CSV Lab")
        self.admin = User.objects.create_user(
            username="csv_admin",
            password="pw",
            email="csv@test.sk",
            role="admin",
            lab=self.lab,
        )

    def _csv_file(self, content):
        from django.core.files.uploadedfile import SimpleUploadedFile

        return SimpleUploadedFile(
            "items.csv", content.encode("utf-8"), content_type="text/csv"
        )

    def test_import_valid_csv(self):
        self.client.force_authenticate(user=self.admin)
        csv_content = "name,sku,quantity,unit,category\nZákladné jehly,SKU-001,50,pcs,consumable\nVosk,SKU-002,20,g,\n"
        resp = self.client.post(
            "/api/inventory/warehouse/import-csv/",
            {"file": self._csv_file(csv_content)},
            format="multipart",
        )
        self.assertEqual(resp.status_code, 201)
        self.assertEqual(resp.data["imported"], 2)
        self.assertEqual(resp.data["skipped"], 0)
        self.assertEqual(WarehouseItem.objects.filter(lab=self.lab).count(), 2)

    def test_import_skips_invalid_rows(self):
        self.client.force_authenticate(user=self.admin)
        csv_content = "name,quantity\nGood Item,10\n,20\n"
        resp = self.client.post(
            "/api/inventory/warehouse/import-csv/",
            {"file": self._csv_file(csv_content)},
            format="multipart",
        )
        self.assertEqual(resp.status_code, 201)
        self.assertEqual(resp.data["imported"], 1)
        self.assertEqual(resp.data["skipped"], 1)

    def test_import_no_file_returns_400(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.post(
            "/api/inventory/warehouse/import-csv/", {}, format="multipart"
        )
        self.assertEqual(resp.status_code, 400)

    def test_unauthenticated_denied(self):
        csv_content = "name\nItem\n"
        resp = self.client.post(
            "/api/inventory/warehouse/import-csv/",
            {"file": self._csv_file(csv_content)},
            format="multipart",
        )
        self.assertEqual(resp.status_code, 401)


class LowStockNotificationTests(APITestCase):
    def setUp(self):
        self.lab = Lab.objects.create(name="Stock Notif Lab")
        self.admin = User.objects.create_user(
            username="stock_admin",
            password="pw",
            email="stock@test.sk",
            role="admin",
            lab=self.lab,
        )

    def test_create_below_threshold_generates_notification(self):
        from apps.core.models import Notification

        self.client.force_authenticate(user=self.admin)
        resp = self.client.post(
            "/api/inventory/warehouse/",
            {"name": "Akrylát", "quantity": 5, "min_threshold": 10},
            format="json",
        )
        self.assertEqual(resp.status_code, 201)
        item_id = resp.data["id"]
        notif = Notification.objects.filter(
            lab=self.lab, type="stock", recipient=self.admin
        ).first()
        self.assertIsNotNone(notif)
        self.assertIn("Akrylát", notif.title)
        self.assertEqual(notif.url, f"/inventory/{item_id}")

    def test_create_above_threshold_no_notification(self):
        from apps.core.models import Notification

        self.client.force_authenticate(user=self.admin)
        self.client.post(
            "/api/inventory/warehouse/",
            {"name": "Composite", "quantity": 50, "min_threshold": 10},
            format="json",
        )
        self.assertEqual(
            Notification.objects.filter(lab=self.lab, type="stock").count(), 0
        )

    def test_update_to_low_stock_creates_notification(self):
        from apps.core.models import Notification

        self.client.force_authenticate(user=self.admin)
        create_resp = self.client.post(
            "/api/inventory/warehouse/",
            {"name": "Zircón", "quantity": 50, "min_threshold": 10},
            format="json",
        )
        item_id = create_resp.data["id"]
        patch_resp = self.client.patch(
            f"/api/inventory/warehouse/{item_id}/",
            {"quantity": 3},
            format="json",
        )
        self.assertEqual(patch_resp.status_code, 200)
        notif = Notification.objects.filter(
            lab=self.lab, type="stock", recipient=self.admin
        ).first()
        self.assertIsNotNone(notif)

    def test_dedup_does_not_create_second_notification_while_unread(self):
        from apps.core.models import Notification

        self.client.force_authenticate(user=self.admin)
        create_resp = self.client.post(
            "/api/inventory/warehouse/",
            {"name": "Wax", "quantity": 3, "min_threshold": 10},
            format="json",
        )
        item_id = create_resp.data["id"]
        # Update again still below threshold
        self.client.patch(
            f"/api/inventory/warehouse/{item_id}/",
            {"quantity": 2},
            format="json",
        )
        self.assertEqual(
            Notification.objects.filter(
                lab=self.lab, type="stock", url=f"/inventory/{item_id}"
            ).count(),
            1,
        )


class InventoryXlsxExportTests(APITestCase):
    def setUp(self):
        self.lab = Lab.objects.create(name="XLSX Export Lab")
        self.admin = User.objects.create_user(
            username="xlsx_inv_admin",
            password="pw",
            email="xlsx_inv@lab.sk",
            role="admin",
            lab=self.lab,
        )
        WarehouseItem.objects.create(
            lab=self.lab,
            name="Zirkón blok",
            sku="ZIR-001",
            quantity=50,
            unit="pcs",
            category="material",
        )
        WarehouseItem.objects.create(
            lab=self.lab,
            name="Separačný lak",
            sku="SEP-002",
            quantity=10,
            unit="ml",
            category="consumable",
        )

    def test_export_xlsx_returns_spreadsheet(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/inventory/warehouse/export/?export_format=xlsx")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])
        self.assertIn(".xlsx", resp["Content-Disposition"])
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        ws = wb.active
        header = [cell.value for cell in ws[1]]
        self.assertIn("name", header)
        self.assertIn("sku", header)
        names = [
            ws.cell(row=r, column=header.index("name") + 1).value
            for r in range(2, ws.max_row + 1)
        ]
        self.assertIn("Zirkón blok", names)
        self.assertIn("Separačný lak", names)

    def test_export_csv_still_works(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/inventory/warehouse/export/")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("text/csv", resp["Content-Type"])
        content = resp.content.decode("utf-8")
        self.assertIn("Zirkón blok", content)

    def test_export_unauthenticated_returns_401(self):
        resp = self.client.get("/api/inventory/warehouse/export/?export_format=xlsx")
        self.assertEqual(resp.status_code, 401)

    @override_settings(EXPORT_MAX_ROWS=1)
    def test_export_enforces_row_limit(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/inventory/warehouse/export/")
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.data["code"], "export_row_limit_exceeded")
        self.assertEqual(str(resp.data["max_rows"]), "1")


class InventoryRoleMatrixTests(RoleMatrixTestMixin, APITestCase):
    """
    Full role matrix coverage for WarehouseItem endpoints.

    Covers: anonymous 401, no_lab, user, technician, admin, superadmin
    for list, create, retrieve, update, and delete.
    """

    def setUp(self):
        self.setup_role_matrix(prefix="inv_rm")
        self.item = WarehouseItem.objects.create(
            lab=self.lab_a,
            name="Matrix Item",
            sku="INVMAT-SKU-001",
            quantity=10,
        )
        self._sku_counter = 1

    def _next_sku(self):
        self._sku_counter += 1
        return f"INVMAT-SKU-{self._sku_counter:04d}"

    def _fresh_item(self):
        return WarehouseItem.objects.create(
            lab=self.lab_a,
            name="Fresh Item",
            sku=self._next_sku(),
            quantity=5,
        )

    # ── List ─────────────────────────────────────────────────────────────────

    def test_warehouse_list_role_matrix(self):
        self.assert_endpoint_matrix(
            "GET",
            "/api/inventory/warehouse/",
            {
                "anonymous": status.HTTP_401_UNAUTHORIZED,
                "no_lab": status.HTTP_200_OK,
                "user": status.HTTP_200_OK,
                "technician": status.HTTP_200_OK,
                "admin": status.HTTP_200_OK,
                "superadmin": status.HTTP_200_OK,
            },
        )

    # ── Create ───────────────────────────────────────────────────────────────

    def test_warehouse_create_role_matrix(self):
        def _payload(role):
            return {
                "lab": self.lab_a.id,
                "name": f"Matrix item {role}",
                "sku": f"INVRM-{role.upper()[:6]}",
                "quantity": 2,
            }

        # anonymous
        resp = self.client.post(
            "/api/inventory/warehouse/",
            _payload("anonymous"),
            format="json",
        )
        self.assertEqual(resp.status_code, status.HTTP_401_UNAUTHORIZED)

        deny_roles = ["no_lab", "user", "technician"]
        for role in deny_roles:
            with self.subTest(role=role):
                self.client.force_authenticate(user=self.role_users[role])
                resp = self.client.post(
                    "/api/inventory/warehouse/",
                    _payload(role),
                    format="json",
                )
                self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)
                self.client.force_authenticate(user=None)

        for role in ["admin", "superadmin"]:
            with self.subTest(role=role):
                self.client.force_authenticate(user=self.role_users[role])
                resp = self.client.post(
                    "/api/inventory/warehouse/",
                    _payload(role),
                    format="json",
                )
                self.assertEqual(resp.status_code, status.HTTP_201_CREATED)
                self.client.force_authenticate(user=None)

    # ── Retrieve ─────────────────────────────────────────────────────────────

    def test_warehouse_retrieve_role_matrix(self):
        self.assert_endpoint_matrix(
            "GET",
            f"/api/inventory/warehouse/{self.item.id}/",
            {
                "anonymous": status.HTTP_401_UNAUTHORIZED,
                "no_lab": status.HTTP_404_NOT_FOUND,
                "user": status.HTTP_200_OK,
                "technician": status.HTTP_200_OK,
                "admin": status.HTTP_200_OK,
                "superadmin": status.HTTP_200_OK,
            },
        )

    # ── Update ───────────────────────────────────────────────────────────────

    def test_warehouse_update_role_matrix(self):
        # no_lab gets 403 (permission check fires before queryset scoping for writes)
        self.assert_endpoint_matrix(
            "PATCH",
            f"/api/inventory/warehouse/{self.item.id}/",
            {
                "anonymous": status.HTTP_401_UNAUTHORIZED,
                "no_lab": status.HTTP_403_FORBIDDEN,
                "user": status.HTTP_403_FORBIDDEN,
                "technician": status.HTTP_403_FORBIDDEN,
                "admin": status.HTTP_200_OK,
                "superadmin": status.HTTP_200_OK,
            },
            data={"quantity": 20},
            format="json",
        )

    # ── Delete ───────────────────────────────────────────────────────────────

    def test_warehouse_delete_role_matrix(self):
        def _delete_matrix(role):
            item = self._fresh_item()
            self.client.force_authenticate(user=self.role_users[role])
            resp = self.client.delete(f"/api/inventory/warehouse/{item.id}/")
            self.client.force_authenticate(user=None)
            return resp

        # anonymous
        anon_item = self._fresh_item()
        resp = self.client.delete(f"/api/inventory/warehouse/{anon_item.id}/")
        self.assertEqual(resp.status_code, status.HTTP_401_UNAUTHORIZED)

        # no_lab gets 403 (permission check fires before queryset scoping for writes)
        expectations = {
            "no_lab": status.HTTP_403_FORBIDDEN,
            "user": status.HTTP_403_FORBIDDEN,
            "technician": status.HTTP_403_FORBIDDEN,
            "admin": status.HTTP_204_NO_CONTENT,
            "superadmin": status.HTTP_204_NO_CONTENT,
        }
        for role, expected in expectations.items():
            with self.subTest(role=role):
                resp = _delete_matrix(role)
                self.assertEqual(
                    resp.status_code, expected, f"DELETE warehouse as {role}"
                )
