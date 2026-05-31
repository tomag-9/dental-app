from django.test import override_settings
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APITestCase

from apps.core.models import Lab, User
from apps.finance.models import PriceList


class PriceListCrudApiTests(APITestCase):
    """Test CRUD operations for PriceList model."""

    def setUp(self):
        self.lab_a = Lab.objects.create(name="Lab A")
        self.lab_b = Lab.objects.create(name="Lab B")

        self.admin_a = User.objects.create_user(
            username="admin_a",
            email="pricelist_admin_a@example.com",
            password="password123",
            role="admin",
            lab=self.lab_a,
        )
        self.admin_b = User.objects.create_user(
            username="admin_b",
            email="pricelist_admin_b@example.com",
            password="password123",
            role="admin",
            lab=self.lab_b,
        )
        self.regular_a = User.objects.create_user(
            username="pricelist_regular_a",
            email="pricelist_regular_a@example.com",
            password="password123",
            role="user",
            lab=self.lab_a,
        )
        self.superadmin = User.objects.create_user(
            username="pricelist_superadmin",
            email="pricelist_superadmin@example.com",
            password="password123",
            role="superadmin",
            is_superuser=True,
        )

    def test_create_price_list_item(self):
        """Test creating a new price list item."""
        self.client.force_authenticate(user=self.admin_a)
        url = reverse("pricelist-list")
        payload = {
            "code": "CROWN123",
            "description": "Test item",
            "price": 12.5,
        }

        response = self.client.post(url, payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["code"], "CROWN123")
        self.assertEqual(response.data["description"], "Test item")
        self.assertEqual(float(response.data["price"]), 12.5)
        self.assertEqual(response.data["lab"], self.lab_a.id)

    def test_superadmin_can_create_price_list_item_for_selected_lab(self):
        self.client.force_authenticate(user=self.superadmin)
        response = self.client.post(
            reverse("pricelist-list"),
            {
                "lab": self.lab_b.id,
                "code": "SUPER-001",
                "description": "Superadmin item",
                "price": 99.0,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["lab"], self.lab_b.id)

    def test_list_price_list_items(self):
        """Test listing price list items scoped to lab."""
        PriceList.objects.create(
            lab=self.lab_a,
            code="ITEM-A1",
            description="Item A1",
            price=10.0,
        )
        PriceList.objects.create(
            lab=self.lab_a,
            code="ITEM-A2",
            description="Item A2",
            price=20.0,
        )
        PriceList.objects.create(
            lab=self.lab_b,
            code="ITEM-B1",
            description="Item B1",
            price=30.0,
        )

        self.client.force_authenticate(user=self.admin_a)
        url = reverse("pricelist-list")

        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data), 2)  # Only lab_a items

    def test_list_price_list_items_supports_bounded_pagination(self):
        for index in range(3):
            PriceList.objects.create(
                lab=self.lab_a,
                code=f"PAGE-{index}",
                description=f"Page item {index}",
                price=10.0,
            )

        self.client.force_authenticate(user=self.admin_a)
        response = self.client.get(reverse("pricelist-list"), {"page_size": 2})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["count"], 3)
        self.assertEqual(len(response.data["results"]), 2)

        page_response = self.client.get(reverse("pricelist-list"), {"page": 1})
        self.assertEqual(page_response.status_code, status.HTTP_200_OK)
        self.assertEqual(page_response.data["count"], 3)
        self.assertEqual(len(page_response.data["results"]), 3)

    def test_superadmin_lists_price_list_items_across_labs(self):
        PriceList.objects.create(
            lab=self.lab_a,
            code="ITEM-A1",
            description="Item A1",
            price=10.0,
        )
        PriceList.objects.create(
            lab=self.lab_b,
            code="ITEM-B1",
            description="Item B1",
            price=30.0,
        )

        self.client.force_authenticate(user=self.superadmin)
        response = self.client.get(reverse("pricelist-list"))

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data), 2)

    def test_duplicate_price_list_code_allowed_across_labs(self):
        PriceList.objects.create(
            lab=self.lab_a,
            code="DUP",
            description="Lab A item",
            price=10.0,
        )

        self.client.force_authenticate(user=self.admin_b)
        response = self.client.post(
            reverse("pricelist-list"),
            {"code": "DUP", "description": "Lab B item", "price": 20.0},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["lab"], self.lab_b.id)

    def test_duplicate_price_list_code_rejected_within_lab(self):
        PriceList.objects.create(
            lab=self.lab_a,
            code="DUP",
            description="Lab A item",
            price=10.0,
        )

        self.client.force_authenticate(user=self.admin_a)
        response = self.client.post(
            reverse("pricelist-list"),
            {"code": "DUP", "description": "Second item", "price": 20.0},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_get_price_list_item(self):
        """Test retrieving a specific price list item."""
        item = PriceList.objects.create(
            lab=self.lab_a,
            code="TEST-001",
            description="Test item",
            price=15.0,
        )

        self.client.force_authenticate(user=self.admin_a)
        url = reverse("pricelist-detail", args=[item.id])

        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["code"], "TEST-001")

    def test_update_price_list_item(self):
        """Test updating a price list item."""
        item = PriceList.objects.create(
            lab=self.lab_a,
            code="UPDATE-001",
            description="Test item",
            price=12.5,
        )

        self.client.force_authenticate(user=self.admin_a)
        url = reverse("pricelist-detail", args=[item.id])
        payload = {
            "code": "UPDATE-001",
            "description": "Updated description",
            "price": 15.0,
        }

        response = self.client.put(url, payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["description"], "Updated description")
        self.assertEqual(float(response.data["price"]), 15.0)

    def test_delete_price_list_item(self):
        """Test deleting a price list item."""
        item = PriceList.objects.create(
            lab=self.lab_a,
            code="DEL-001",
            description="Test item",
            price=10.0,
        )

        self.client.force_authenticate(user=self.admin_a)
        url = reverse("pricelist-detail", args=[item.id])

        response = self.client.delete(url)

        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)

        # Verify deletion
        response = self.client.get(url)
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_cross_lab_access_denied(self):
        """Test that users cannot access price list items from other labs."""
        item_b = PriceList.objects.create(
            lab=self.lab_b,
            code="LAB-B-001",
            description="Lab B item",
            price=5.0,
        )

        self.client.force_authenticate(user=self.admin_a)
        url = reverse("pricelist-detail", args=[item_b.id])

        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_duplicate_price_list_item_creates_copy_in_same_lab(self):
        item = PriceList.objects.create(
            lab=self.lab_a,
            code="CROWN",
            description="Zircon crown",
            price=150.0,
        )

        self.client.force_authenticate(user=self.admin_a)
        response = self.client.post(reverse("pricelist-duplicate", args=[item.id]))

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["lab"], self.lab_a.id)
        self.assertEqual(response.data["code"], "CROWN-COPY")
        self.assertEqual(response.data["description"], "Zircon crown")
        self.assertEqual(float(response.data["price"]), 150.0)

    def test_duplicate_price_list_item_generates_unique_copy_code(self):
        item = PriceList.objects.create(
            lab=self.lab_a,
            code="CROWN",
            description="Zircon crown",
            price=150.0,
        )
        PriceList.objects.create(
            lab=self.lab_a,
            code="CROWN-COPY",
            description="Existing copy",
            price=150.0,
        )

        self.client.force_authenticate(user=self.admin_a)
        response = self.client.post(reverse("pricelist-duplicate", args=[item.id]))

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["code"], "CROWN-COPY-2")

    def test_duplicate_price_list_item_respects_lab_scope(self):
        item_b = PriceList.objects.create(
            lab=self.lab_b,
            code="LAB-B",
            description="Other lab item",
            price=10.0,
        )

        self.client.force_authenticate(user=self.admin_a)
        response = self.client.post(reverse("pricelist-duplicate", args=[item_b.id]))

        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_regular_user_cannot_create_price_list_item(self):
        self.client.force_authenticate(user=self.regular_a)
        response = self.client.post(
            reverse("pricelist-list"),
            {"code": "REG-001", "description": "Regular item", "price": 12.0},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertFalse(PriceList.objects.filter(code="REG-001").exists())

    def test_regular_user_cannot_duplicate_price_list_item(self):
        item = PriceList.objects.create(
            lab=self.lab_a,
            code="REG-DUP",
            description="Regular duplicate",
            price=10.0,
        )

        self.client.force_authenticate(user=self.regular_a)
        response = self.client.post(reverse("pricelist-duplicate", args=[item.id]))

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertFalse(PriceList.objects.filter(code="REG-DUP-COPY").exists())


class ProcedureCatalogTests(APITestCase):
    def setUp(self):
        self.lab = Lab.objects.create(name="Catalog Lab")
        self.admin = User.objects.create_user(
            username="cat_admin",
            email="cat_admin@example.com",
            password="pass",
            role="admin",
            lab=self.lab,
        )
        self.other_lab = Lab.objects.create(name="Other Lab")
        self.other_user = User.objects.create_user(
            username="other_u",
            email="other_u@example.com",
            password="pass",
            role="user",
            lab=self.other_lab,
        )
        PriceList.objects.create(
            lab=self.lab,
            code="C001",
            description="Full crown",
            price="150.00",
            category="crown",
        )
        PriceList.objects.create(
            lab=self.lab,
            code="B001",
            description="3-unit bridge",
            price="400.00",
            category="bridge",
        )
        PriceList.objects.create(lab=self.lab, code="X001", description="Misc", price="50.00", category=None)
        PriceList.objects.create(
            lab=self.other_lab,
            code="C001",
            description="Other crown",
            price="200.00",
            category="crown",
        )

    def test_catalog_returns_own_lab_items_grouped(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/finance/procedure-catalog/")
        self.assertEqual(resp.status_code, 200)
        categories = {g["category"] for g in resp.data}
        self.assertIn("crown", categories)
        self.assertIn("bridge", categories)
        crown_group = next(g for g in resp.data if g["category"] == "crown")
        self.assertEqual(len(crown_group["items"]), 1)
        self.assertEqual(crown_group["items"][0]["code"], "C001")

    def test_catalog_excludes_other_lab_items(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/finance/procedure-catalog/")
        all_ids = [i["id"] for g in resp.data for i in g["items"]]
        other_ids = list(PriceList.objects.filter(lab=self.other_lab).values_list("id", flat=True))
        for oid in other_ids:
            self.assertNotIn(oid, all_ids)

    def test_uncategorized_items_grouped_separately(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/finance/procedure-catalog/")
        uncat_group = next((g for g in resp.data if g["category"] is None), None)
        self.assertIsNotNone(uncat_group)
        self.assertEqual(uncat_group["label"], "Uncategorized")
        self.assertEqual(len(uncat_group["items"]), 1)

    def test_unauthenticated_returns_401(self):
        resp = self.client.get("/api/finance/procedure-catalog/")
        self.assertEqual(resp.status_code, 401)


class PriceListExportCsvTests(APITestCase):
    def setUp(self):
        self.lab = Lab.objects.create(name="PL Export Lab")
        self.admin = User.objects.create_user(
            username="pl_export_admin",
            password="pw",
            email="pl@lab.sk",
            role="admin",
            lab=self.lab,
        )
        PriceList.objects.create(
            lab=self.lab,
            code="KOR-001",
            description="Korunka zirkónová",
            price="280.00",
            category="crown",
        )
        PriceList.objects.create(
            lab=self.lab,
            code="MOS-002",
            description="Mostík 3-členný",
            price="650.00",
            category="bridge",
        )

    def test_export_returns_csv_with_header(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/finance/price-list/export/")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("text/csv", resp["Content-Type"])
        content = (
            b"".join(resp.streaming_content).decode() if hasattr(resp, "streaming_content") else resp.content.decode()
        )
        self.assertIn("code,description,price", content)

    def test_export_contains_all_lab_items(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/finance/price-list/export/")
        content = (
            b"".join(resp.streaming_content).decode() if hasattr(resp, "streaming_content") else resp.content.decode()
        )
        rows = [r for r in content.strip().split("\n") if r]
        self.assertEqual(len(rows), 3)  # header + 2 items
        self.assertIn("KOR-001", content)
        self.assertIn("MOS-002", content)

    def test_export_tenant_scoped(self):
        other_lab = Lab.objects.create(name="Other PL Lab")
        other_user = User.objects.create_user(
            username="other_pl_admin",
            password="pw",
            email="other_pl@lab.sk",
            role="admin",
            lab=other_lab,
        )
        self.client.force_authenticate(user=other_user)
        resp = self.client.get("/api/finance/price-list/export/")
        content = (
            b"".join(resp.streaming_content).decode() if hasattr(resp, "streaming_content") else resp.content.decode()
        )
        self.assertNotIn("KOR-001", content)

    def test_export_unauthenticated_denied(self):
        resp = self.client.get("/api/finance/price-list/export/")
        self.assertEqual(resp.status_code, 401)

    def test_export_xlsx_returns_spreadsheet(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/finance/price-list/export/?export_format=xlsx")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])
        self.assertIn(".xlsx", resp["Content-Disposition"])
        from io import BytesIO

        import openpyxl

        wb = openpyxl.load_workbook(BytesIO(resp.content))
        ws = wb.active
        header = [cell.value for cell in ws[1]]
        self.assertIn("code", header)
        codes = [ws.cell(row=r, column=header.index("code") + 1).value for r in range(2, ws.max_row + 1)]
        self.assertIn("KOR-001", codes)
        self.assertIn("MOS-002", codes)

    @override_settings(EXPORT_MAX_ROWS=1)
    def test_export_enforces_row_limit(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/finance/price-list/export/")
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.data["code"], "export_row_limit_exceeded")
        self.assertEqual(str(resp.data["max_rows"]), "1")
