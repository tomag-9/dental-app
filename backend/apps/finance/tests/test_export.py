from django.test import override_settings
from django.utils import timezone
from rest_framework.test import APITestCase

from apps.core.models import Lab, User
from apps.crm.models import Clinic
from apps.finance.models import Invoice


class InvoiceCSVExportTests(APITestCase):
    def setUp(self):
        self.lab = Lab.objects.create(name="Export Lab")
        self.other_lab = Lab.objects.create(name="Other Lab")
        self.clinic = Clinic.objects.create(lab=self.lab, name="Klinika A")
        self.other_clinic = Clinic.objects.create(lab=self.other_lab, name="Other Clinic")
        self.admin = User.objects.create_user(
            username="export_admin",
            email="export_admin@example.com",
            password="pass",
            role="admin",
            lab=self.lab,
        )
        now = timezone.now()
        Invoice.objects.create(
            lab=self.lab,
            clinic=self.clinic,
            number="EXP-001",
            status="paid",
            total_amount="100.00",
            issued_at=now,
            paid_at=now,
        )
        Invoice.objects.create(
            lab=self.lab,
            clinic=self.clinic,
            number="EXP-002",
            status="issued",
            total_amount="200.00",
            issued_at=now,
        )
        Invoice.objects.create(
            lab=self.other_lab,
            clinic=self.other_clinic,
            number="OTHER-001",
            status="paid",
            total_amount="999.00",
        )

    def test_export_returns_csv(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/finance/invoices/export/")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("text/csv", resp["Content-Type"])
        content = resp.content.decode("utf-8")
        self.assertIn("Číslo,Stav,Klinika", content)
        self.assertIn("Zaplatená", content)
        self.assertIn("EXP-001", content)
        self.assertIn("EXP-002", content)

    def test_export_excludes_other_lab(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/finance/invoices/export/")
        content = resp.content.decode("utf-8")
        self.assertNotIn("OTHER-001", content)

    def test_export_status_filter(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/finance/invoices/export/?status=paid")
        content = resp.content.decode("utf-8")
        self.assertIn("EXP-001", content)
        self.assertNotIn("EXP-002", content)

    def test_export_unauthenticated_returns_401(self):
        resp = self.client.get("/api/finance/invoices/export/")
        self.assertEqual(resp.status_code, 401)

    def test_export_xlsx_returns_spreadsheet(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/finance/invoices/export/?export_format=xlsx")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])
        self.assertIn(".xlsx", resp["Content-Disposition"])
        from io import BytesIO

        import openpyxl

        wb = openpyxl.load_workbook(BytesIO(resp.content))
        ws = wb.active
        header = [cell.value for cell in ws[1]]
        self.assertIn("Číslo", header)
        numbers = [ws.cell(row=r, column=header.index("Číslo") + 1).value for r in range(2, ws.max_row + 1)]
        self.assertIn("EXP-001", numbers)
        self.assertIn("EXP-002", numbers)

    @override_settings(EXPORT_MAX_ROWS=1)
    def test_export_enforces_row_limit(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/finance/invoices/export/")
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.data["code"], "export_row_limit_exceeded")
        self.assertEqual(str(resp.data["max_rows"]), "1")
