from django.test import override_settings
from rest_framework.test import APITestCase

from apps.core.models import Lab, User
from apps.crm.models import Clinic, Patient
from apps.jobs.models import Job


class JobExportCsvTests(APITestCase):
    def setUp(self):
        self.lab = Lab.objects.create(name="Export Lab")
        self.admin = User.objects.create_user(
            username="export_admin",
            password="pw",
            email="export@lab.sk",
            role="admin",
            lab=self.lab,
        )
        self.patient = Patient.objects.create(
            first_name="Ján", last_name="Testovský", lab=self.lab
        )
        self.clinic = Clinic.objects.create(name="Klinika Export", lab=self.lab)
        Job.objects.create(
            patient=self.patient,
            clinic=self.clinic,
            lab=self.lab,
            status="completed",
            price="250.00",
        )
        Job.objects.create(
            patient=self.patient,
            clinic=self.clinic,
            lab=self.lab,
            status="new",
            price="100.00",
        )

    def test_export_returns_csv_with_header(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/jobs/jobs/export/")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("text/csv", resp["Content-Type"])
        content = (
            b"".join(resp.streaming_content).decode()
            if hasattr(resp, "streaming_content")
            else resp.content.decode()
        )
        self.assertIn("id,status,patient", content)

    def test_export_contains_all_lab_jobs(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/jobs/jobs/export/")
        content = (
            b"".join(resp.streaming_content).decode()
            if hasattr(resp, "streaming_content")
            else resp.content.decode()
        )
        rows = [r for r in content.strip().split("\n") if r]
        self.assertEqual(len(rows), 3)  # header + 2 jobs

    def test_export_respects_status_filter(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/jobs/jobs/export/?status=completed")
        content = (
            b"".join(resp.streaming_content).decode()
            if hasattr(resp, "streaming_content")
            else resp.content.decode()
        )
        rows = [r for r in content.strip().split("\n") if r]
        self.assertEqual(len(rows), 2)  # header + 1 completed job

    def test_export_tenant_scoped(self):
        other_lab = Lab.objects.create(name="Other Lab")
        other_user = User.objects.create_user(
            username="other_export",
            password="pw",
            email="other@lab.sk",
            role="admin",
            lab=other_lab,
        )
        self.client.force_authenticate(user=other_user)
        resp = self.client.get("/api/jobs/jobs/export/")
        content = (
            b"".join(resp.streaming_content).decode()
            if hasattr(resp, "streaming_content")
            else resp.content.decode()
        )
        rows = [r for r in content.strip().split("\n") if r]
        self.assertEqual(len(rows), 1)  # header only, no jobs from other lab

    def test_export_unauthenticated_denied(self):
        resp = self.client.get("/api/jobs/jobs/export/")
        self.assertEqual(resp.status_code, 401)

    def test_export_xlsx_returns_spreadsheet(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/jobs/jobs/export/?export_format=xlsx")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])
        self.assertIn(".xlsx", resp["Content-Disposition"])
        from io import BytesIO

        import openpyxl

        wb = openpyxl.load_workbook(BytesIO(resp.content))
        ws = wb.active
        header = [cell.value for cell in ws[1]]
        self.assertIn("id", header)
        self.assertIn("status", header)
        self.assertGreaterEqual(ws.max_row, 2)

    @override_settings(EXPORT_MAX_ROWS=1)
    def test_export_enforces_row_limit(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/jobs/jobs/export/")
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.data["code"], "export_row_limit_exceeded")
        self.assertEqual(str(resp.data["max_rows"]), "1")
